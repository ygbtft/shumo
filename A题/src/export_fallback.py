import json
from pathlib import Path
import re
import sys
import zipfile

from openpyxl import Workbook, load_workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Alignment, Font, PatternFill


PROJECT = Path(__file__).resolve().parents[1]


def compact_dense_workbook(destination):
    temporary = destination.with_suffix(".compact.xlsx")
    with zipfile.ZipFile(destination) as original, zipfile.ZipFile(temporary, "w", zipfile.ZIP_DEFLATED, compresslevel=9) as compact:
        for item in original.infolist():
            content = original.read(item.filename)
            if item.filename in ("xl/worksheets/sheet1.xml", "xl/worksheets/sheet2.xml"):
                content = re.sub(rb' r="[A-Z]+[0-9]+"', b"", content)
            compact.writestr(item.filename, content)
    temporary.replace(destination)


def export(entry):
    workbook = Workbook(write_only=True)
    workbook.properties.creator = ""
    workbook.properties.lastModifiedBy = ""
    workbook.properties.title = "药材烘干计算结果"
    for definition in entry["sheets"]:
        sheet = workbook.create_sheet(definition["name"])
        sheet.freeze_panes = "B2"
        sheet.sheet_view.showGridLines = False
        sheet.column_dimensions["A"].width = 25
        for column_index in range(1, definition["columns"]):
            from openpyxl.utils import get_column_letter
            sheet.column_dimensions[get_column_letter(column_index + 1)].width = 12
        rows = json.loads((PROJECT / definition["data"]).read_text(encoding="utf-8"))
        header = []
        for value in rows[0]:
            cell = WriteOnlyCell(sheet, value)
            cell.font = Font(name="Microsoft YaHei", size=10, bold=True)
            cell.fill = PatternFill("solid", fgColor="E7ECF2")
            cell.alignment = Alignment(horizontal="center", vertical="center")
            header.append(cell)
        sheet.row_dimensions[1].height = 28
        sheet.append(header)
        for row in rows[1:]:
            cells = []
            for column_index, value in enumerate(row):
                if value is None:
                    cells.append(None)
                    continue
                cell = WriteOnlyCell(sheet, value)
                cell.number_format = "0" if column_index == 0 else "0.0000"
                cells.append(cell)
            sheet.append(cells)
        print(entry["filename"], definition["name"], len(rows), flush=True)
        del rows
    notes = workbook.create_sheet("说明")
    notes.column_dimensions["A"].width = 18
    notes.column_dimensions["B"].width = 85
    notes.append(["项目", "说明"])
    notes.append(["输入", "A题附件1.xlsx、附件2.xlsx；物性参数见题目附录2至4"])
    notes.append(["方法", "积分势通量、加密有限体积网格、BDF；外部程序计算，数值保留四位小数"])
    notes.append(["单位", "时间s；距离cm；温度摄氏度；含水率kg/kg干基"])
    notes.append(["空白", "固定距离超过当时药材半径，表示不在药材内部，不代表含水率为零"])
    notes.append(["输出范围", "result2为完整干燥过程逐秒结果；result2_3h是前三小时的便捷节选"])
    for row_number in range(1, 8):
        notes.row_dimensions[row_number].height = 32
    destination = PROJECT / "results/workbooks" / entry["filename"]
    destination.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(destination)
    if entry["filename"] in ("result2.xlsx", "result2_3h.xlsx"):
        compact_dense_workbook(destination)
    checked = load_workbook(destination, read_only=True, data_only=True)
    for definition in entry["sheets"]:
        sheet = checked[definition["name"]]
        count = 0
        last = None
        for row in sheet.iter_rows(values_only=True):
            count += 1
            last = row
            if count > 1:
                assert isinstance(row[0], (int, float))
                assert all(value is None or isinstance(value, (int, float)) for value in row)
        assert count == definition["rows"]
        print("VERIFIED", entry["filename"], definition["name"], count, "last_time", last[0], flush=True)
    checked.close()


def main():
    selected = sys.argv[1] if len(sys.argv) > 1 else "core"
    manifest = json.loads((PROJECT / "cache/export_manifest.json").read_text(encoding="utf-8"))
    for entry in manifest:
        if selected == "all" or (selected == "core" and "_3h" not in entry["filename"]) or selected == entry["filename"]:
            export(entry)


if __name__ == "__main__":
    main()
