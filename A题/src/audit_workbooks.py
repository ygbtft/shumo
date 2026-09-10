from collections import deque
from html import escape
import json
import os
from pathlib import Path

from openpyxl import load_workbook
from reportlab.lib import colors
from reportlab.lib.pagesizes import A3, landscape
from reportlab.lib.styles import ParagraphStyle
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle


PROJECT = Path(__file__).resolve().parents[1]


def main():
    font_path = Path(os.environ.get("CJK_FONT_REGULAR", str(Path(os.environ.get("WINDIR", "C:/Windows")) / "Fonts/simsun.ttc")))
    pdfmetrics.registerFont(TTFont("CJK", str(font_path)))
    appearance = ParagraphStyle("cell", fontName="CJK", fontSize=9, leading=13, wordWrap="CJK")
    title_style = ParagraphStyle("title", fontName="CJK", fontSize=14, leading=21, spaceAfter=12)
    manifest = json.loads((PROJECT / "cache/export_manifest.json").read_text(encoding="utf-8"))
    records, story = [], []
    for entry in manifest:
        destination = PROJECT / "results/workbooks" / entry["filename"]
        workbook = load_workbook(destination, read_only=True, data_only=False)
        assert not workbook.properties.creator and not workbook.properties.lastModifiedBy
        sheets = []
        for definition in entry["sheets"]:
            expected = json.loads((PROJECT / definition["data"]).read_text(encoding="utf-8"))
            sheet = workbook[definition["name"]]
            initial, final = [], deque(maxlen=4)
            count, maximum_difference = 0, 0.0
            for count, row in enumerate(sheet.iter_rows(), 1):
                values = [cell.value for cell in row]
                reference = expected[count - 1]
                assert len(values) == len(reference), (entry["filename"], definition["name"], count)
                for column, (cell, wanted) in enumerate(zip(row, reference)):
                    assert cell.data_type != "f" and cell.data_type != "e"
                    if wanted is None or isinstance(wanted, str):
                        assert cell.value == wanted
                    else:
                        assert isinstance(cell.value, (int, float))
                        maximum_difference = max(maximum_difference, abs(cell.value - wanted))
                    if count > 1 and cell.value is not None:
                        assert cell.number_format == ("0" if column == 0 else "0.0000")
                if count <= 9:
                    initial.append(values)
                if count > 1:
                    final.append(values)
            assert count == definition["rows"]
            assert maximum_difference < 1e-12
            sheets.append({"sheet": definition["name"], "rows_with_header": count, "columns": definition["columns"],
                           "last_time_s": final[-1][0], "maximum_roundtrip_difference": maximum_difference})
            selected = [*initial, ["...", *[""] * (definition["columns"] - 1)], *final]
            story.append(Paragraph(escape(entry["filename"] + " / " + definition["name"]), title_style))
            story.append(Paragraph("保存文件回读预览：首8条与末4条记录。该预览不是Excel渲染引擎截图。", appearance))
            story.append(Spacer(1, 12))
            cells = [[Paragraph(escape("" if value is None else str(value) if row_index == 0 or column == 0 or isinstance(value, str) else f"{value:.4f}"), appearance)
                      for column, value in enumerate(row)] for row_index, row in enumerate(selected)]
            columns = definition["columns"]
            total_width = landscape(A3)[0] - 72
            widths = [116, *[(total_width - 116) / (columns - 1)] * (columns - 1)] if columns > 2 else [250, 400]
            table = Table(cells, colWidths=widths, repeatRows=1)
            table.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#E7ECF2")),
                                       ("VALIGN", (0, 0), (-1, -1), "MIDDLE"), ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
                                       ("TOPPADDING", (0, 0), (-1, -1), 6), ("LINEBELOW", (0, 0), (-1, 0), 0.5, colors.grey)]))
            story.extend([table, PageBreak()])
            del expected
        notes = list(workbook["说明"].values)
        story.append(Paragraph(escape(entry["filename"] + " / 说明"), title_style))
        note_table = Table([[Paragraph(escape(str(value or "")), appearance) for value in row] for row in notes], colWidths=[150, 720])
        note_table.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "TOP"), ("BOTTOMPADDING", (0, 0), (-1, -1), 12)]))
        story.extend([note_table, PageBreak()])
        workbook.close()
        records.append({"workbook": entry["filename"], "bytes": destination.stat().st_size, "sheets": sheets})
        print("逐单元格核验", entry["filename"], sheets, flush=True)
    preview = PROJECT / "tmp/workbook_previews/工作簿回读预览.pdf"
    preview.parent.mkdir(parents=True, exist_ok=True)
    SimpleDocTemplate(str(preview), pagesize=landscape(A3), leftMargin=36, rightMargin=36, topMargin=36, bottomMargin=36,
                      author="", title="工作簿回读预览").build(story[:-1])
    (PROJECT / "results/workbook_audit.json").write_text(json.dumps(records, ensure_ascii=False, indent=2), encoding="utf-8")


if __name__ == "__main__":
    main()
